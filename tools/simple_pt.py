"""simple_pt — "Simple is best" 양식의 발표 자료 (pptx) 를 만드는 작은 모듈 (python-pptx).

디자인 값 (자리·글꼴·크기·색) 이 이 파일 한 곳에 있어서, 누가 어떤 AI 로 만들어도 같은 모양이 나온다.
규칙은 rules/ — design (모양), writing (글), structure (덱 구성), checklist (검수).

쓰는 법:
    from simple_pt import Deck
    d = Deck("Demo")
    d.cover("Short Topic:\\nWhat It Does", subtitle="한 줄 부제", kicker="Team Seminar",
            name="Your Name", affiliation="Dept. · Lab", date="YYYY-MM-DD")
    d.panels("Results", "이 장의 결론 한 문장 (음슴체)",
             panels=[{"label": "Gain", "image": "figures/gain.png", "caption": "무엇을 봐야 하는지 한 줄"}],
             key="빨간 결론 한 줄", notes="말할 대본")
    d.save("deck.pptx")                # 규칙에 어긋나 보이는 곳은 경고로 알린다

장 종류 (레이아웃 이름은 rules/design.md 7 절):
    cover           표지
    panels          그림 1–3 칸 (칸마다 이름표·그림·캡션) + 빨간 줄            F
    figure_bullets  왼쪽 그림 + 오른쪽 불릿 3–5 개 + 빨간 줄                   A
    cards           카드 2–3 개 (소제목 + 불릿) + 빨간 줄                      E·G
    equation        회색 수식 박스 (식 + where 한 문장) + 아래 그림·카드        B
    table           표 + 빨간 줄                                               C
    summary         얇은 선 아래 3 단 열 + 회색 마무리 문장                    H
    appendix_index  Appendix 첫 장 — 질문 | Appendix | 장 (장 번호는 save 때 채운다)
appendix_index 뒤의 장은 쪽 번호가 A1, A2 … 가 된다.

글 표시: **굵게**, V_{IN} 아래 첨자, x^{2} 위 첨자. 줄바꿈은 "\\n", 불릿은 리스트.
"""
import datetime
import math
import re
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ── 디자인 값 — 바꾸면 rules/design.md 도 같이 고친다 ─────────────────────────
W, H = 1920, 1080                  # 캔버스 px (= 13.333 × 7.5 in, 1 pt = 2 px)
MX, TOP, BOTTOM = 128, 72, 920     # 좌우 여백, 제목 위 끝, 본문 아래 끝
CW = W - 2 * MX                    # 본문 폭 1664
PAGE_Y = H - 64                    # 쪽 번호·표지 날짜의 아래 끝

LATIN, KR = "Arial", "Noto Sans KR"
INK, TEXT, GREY, MUTED, RED = "111111", "1F1F1F", "4A4A4A", "767676", "B4161E"
CARD_FILL, CARD_LINE, HEAD_FILL, GRID, RULE = "F5F5F3", "E5E5E3", "EDEEE9", "CFCFCF", "D9D9D9"

SIZE = {  # 글자 크기 px (pt 의 2 배)
    "title": 52, "subtitle": 40, "label": 28, "body": 32, "card_title": 32, "card_body": 28,
    "table": 28, "caption": 24, "key": 34, "page": 24, "notes": 24,
    "cover_title": 72, "cover_kicker": 24, "cover_name": 32,
}
LINE = {  # 줄 높이 — 글자 크기의 배수 (CSS line-height 와 같은 뜻, pptx 에는 정확한 pt 로 넣는다)
    "title": 1.15, "subtitle": 1.3, "label": 1.3, "body": 1.45, "card": 1.4, "table": 1.3,
    "caption": 1.4, "key": 1.3, "page": 1.2, "cover_title": 1.12, "cover": 1.3,
}
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # PowerPoint "스타일 없음, 눈금 없음"

HANGUL = re.compile(r"[ᄀ-ᇿ㄰-㆏가-힣]")
MARK = re.compile(r"\*\*(.+?)\*\*|_\{(.+?)\}|\^\{(.+?)\}")
SENTENCE_END = re.compile(r"(다|요)[.!]?$")                   # "~다." 로 끝나는 문장형


def px(v):
    return Emu(round(v * 6350))                                # 1 px = 6350 EMU


# ── 글 ──────────────────────────────────────────────────────────────────────
def _plain(text):
    return MARK.sub(lambda m: next(g for g in m.groups() if g is not None), str(text))


def _runs(text):
    """글 표시를 run 조각으로 — [(글, 굵게, baseline)]."""
    out, pos = [], 0
    for m in MARK.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, 0))
        bold, sub, sup = m.groups()
        if bold is not None:
            out.append((bold, True, 0))
        elif sub is not None:
            out.append((sub, False, -25000))
        else:
            out.append((sup, False, 30000))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, 0))
    return out or [("", False, 0)]


def _width(text, size, bold=False):
    """한 줄 글의 폭 어림 (px) — 한글·한자 1.0 em, 영문 대문자 0.68, 그 밖 0.55, 빈칸 0.28."""
    w = 0.0
    for ch in _plain(text):
        if HANGUL.match(ch) or ord(ch) >= 0x2E80:
            w += 1.0
        elif ch == " ":
            w += 0.28
        elif ch.isupper() or ch in "@%":
            w += 0.68
        else:
            w += 0.55
    return w * size * (1.07 if bold else 1.0)


def _lines(text, size, width, bold=False):
    """줄바꿈까지 어림한 줄 수."""
    if not text:
        return 0
    paras = text if isinstance(text, (list, tuple)) else str(text).split("\n")
    return sum(max(1, math.ceil(_width(p, size, bold) / width)) for p in paras)


def _font(run, size, color, bold=False, baseline=0, caps=False, spacing=0):
    """run 의 글꼴 — 영문 Arial, 한글 (ea) Noto Sans KR, 언어는 한글이 있으면 ko-KR (줄바꿈이 낱말 단위가 된다)."""
    f = run.font
    f.size = Pt(size / 2)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = LATIN
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
    ea, cs = etree.Element(qn("a:ea")), etree.Element(qn("a:cs"))
    ea.set("typeface", KR)
    cs.set("typeface", LATIN)
    latin = rpr.find(qn("a:latin"))
    latin.addnext(ea)
    ea.addnext(cs)
    rpr.set("lang", "ko-KR" if HANGUL.search(run.text) else "en-US")
    if baseline:
        rpr.set("baseline", str(baseline))
    if caps:
        rpr.set("cap", "all")
    if spacing:
        rpr.set("spc", str(spacing))


def _bullet(p, size):
    ppr = p._p.get_or_add_pPr()
    ind = int(px(size * 0.85))
    ppr.set("marL", str(ind))
    ppr.set("indent", str(-ind))
    bu_font = etree.SubElement(ppr, qn("a:buFont"))
    bu_font.set("typeface", LATIN)
    bu_char = etree.SubElement(ppr, qn("a:buChar"))
    bu_char.set("char", "•")


def _fill_frame(tf, text, size, color, bold=False, line=1.3, align=PP_ALIGN.LEFT, bullets=False,
                gap_after=0, caps=False, spacing=0):
    paras = list(text) if isinstance(text, (list, tuple)) else str(text).split("\n")
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(size / 2 * line)
        if gap_after and i < len(paras) - 1:
            p.space_after = Pt(gap_after / 2)
        if bullets:
            _bullet(p, size)
        for part, b, base in _runs(str(para)):
            r = p.add_run()
            r.text = part
            _font(r, size, color, bold or b, base, caps, spacing)


def _textbox(slide, x, y, w, h, text, size, color, bold=False, line=1.3, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, bullets=False, gap_after=0, caps=False, spacing=0):
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    _fill_frame(tf, text, size, color, bold, line, align, bullets, gap_after, caps, spacing)
    return box


# ── 도형·그림·표 ─────────────────────────────────────────────────────────────
def _plain_shape(shape):
    """테마 스타일 (그림자·테마 색) 을 떼어 낸다 — 채움·선은 직접 준다."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _rect(slide, x, y, w, h, fill, line=None, radius=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 px(x), px(y), px(w), px(h))
    _plain_shape(shp)
    if radius:
        shp.adjustments[0] = radius / min(w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    return shp


def _rule(slide, x1, y, x2, color=RULE):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y), px(x2), px(y))
    _plain_shape(ln)
    ln.line.color.rgb = RGBColor.from_string(color)
    ln.line.width = Pt(0.75)


def _fixed_zip_dates(path):
    """pptx (zip) 안 파일의 날짜를 고정한다 — 만든 날이 파일에 남지 않고, 같은 입력이면 같은 파일이 된다."""
    with zipfile.ZipFile(path) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            zi = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zout.writestr(zi, data)


def _image_size(path):
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"그림이 없다: {path}")
    with Image.open(path) as im:
        return im.size


def _equation_png(eq):
    """식 그림 — PNG 경로면 그대로, LaTeX 글이면 tools/equations.py 로 임시 폴더에 만든다."""
    if isinstance(eq, (str, Path)) and str(eq).lower().endswith(".png"):
        if not Path(eq).is_file():
            raise FileNotFoundError(f"식 그림이 없다: {eq}")
        return Path(eq)
    import hashlib
    import sys
    import tempfile
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from equations import render
    name = hashlib.sha1(repr(eq).encode("utf-8")).hexdigest()[:12]
    return render(eq, Path(tempfile.gettempdir()) / "simple_pt_eq" / name)


def _picture(slide, path, x, y, w, h, alt=""):
    pic = slide.shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    pic._element.nvPicPr.cNvPr.set("descr", alt or Path(path).stem.replace("_", " "))
    return pic


def _cell(cell, text, size, bold, fill):
    cell.margin_left = cell.margin_right = px(16)
    cell.margin_top = cell.margin_bottom = px(10)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    tf = cell.text_frame
    tf.word_wrap = True
    _fill_frame(tf, text, size, TEXT, bold, LINE["table"])
    tcpr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):   # 칸 테두리 — 채움보다 앞에 둔다
        old = tcpr.find(qn(tag))
        if old is not None:
            tcpr.remove(old)
        ln = etree.Element(qn(tag), w="9525", cap="flat", cmpd="sng", algn="ctr")
        clr = etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"))
        clr.set("val", GRID)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcpr.insert(i, ln)


def _table(slide, x, y, w, header, rows, widths=None, size=SIZE["table"]):
    """표 — 머리 줄은 굵게·연회색, 격자는 회색 0.75 pt. 돌려주는 것: python-pptx 표, 높이 (px)."""
    ncol = len(header)
    widths = widths or [1] * ncol
    col_w = [w * v / sum(widths) for v in widths]
    line_h = size * LINE["table"]
    grid = [list(map(str, header))] + [list(map(str, r)) for r in rows]
    heights = [20 + line_h * max(1, max(_lines(c, size, cw - 32, bold=(r == 0)) for c, cw in zip(cells, col_w)))
               for r, cells in enumerate(grid)]
    frame = slide.shapes.add_table(len(grid), ncol, px(x), px(y), px(w), px(sum(heights)))
    tbl = frame.table
    tblpr = tbl._tbl.tblPr
    tblpr.set("firstRow", "0")
    tblpr.set("bandRow", "0")
    sid = tblpr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblpr, qn("a:tableStyleId"))
    sid.text = NO_TABLE_STYLE
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = px(cw)
    for r, hgt in enumerate(heights):
        tbl.rows[r].height = px(hgt)
    for r, cells in enumerate(grid):
        for c, val in enumerate(cells):
            _cell(tbl.cell(r, c), val, size, bold=(r == 0), fill=HEAD_FILL if r == 0 else "FFFFFF")
    return tbl, sum(heights)


# ── 덱 ──────────────────────────────────────────────────────────────────────
class Deck:
    """장을 순서대로 더하고 save() 로 쓴다. 장마다 notes 에 말할 대본을 넣는다."""

    def __init__(self, title=""):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = px(W), px(H)
        self.title = title
        self.warnings = []
        self._count = 0                 # 장 수
        self._appendix = None           # appendix_index 뒤로 몇 번째 A 장인가
        self._labels = {}               # 쪽 번호 글 → 장 번호
        self._page = {}                 # slide_id → 쪽 번호 글
        self._index_cells = []          # (표 칸, 가리키는 쪽 번호 글)

    # 공통 ------------------------------------------------------------------
    def _warn(self, page, msg):
        self.warnings.append(f"[{page}] {msg}")

    def _new_slide(self):
        self._count += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])     # 빈 레이아웃

    def _start(self, title, subtitle, key, page=None):
        s = self._new_slide()
        if page is None:
            if self._appendix is None:
                page = str(self._count)
            else:
                self._appendix += 1
                page = f"A{self._appendix}"
        self._labels[page] = self._count
        self._page[s.slide_id] = page
        self._check_text(page, title, subtitle, key)
        th = SIZE["title"] * LINE["title"]
        _textbox(s, MX, TOP, CW, th, title, SIZE["title"], INK, bold=True, line=LINE["title"])
        top = TOP + th + 12
        if subtitle:
            n = _lines(subtitle, SIZE["subtitle"], CW)
            sh = n * SIZE["subtitle"] * LINE["subtitle"]
            _textbox(s, MX, top, CW, sh, subtitle, SIZE["subtitle"], GREY, line=LINE["subtitle"])
            top += sh
        top += 18 + 40
        bottom = BOTTOM
        if key:
            kh = _lines(key, SIZE["key"], CW, bold=True) * SIZE["key"] * LINE["key"]
            _textbox(s, MX, BOTTOM - kh, CW, kh, key, SIZE["key"], RED, bold=True, line=LINE["key"])
            bottom = BOTTOM - kh - 40
        return s, top, bottom

    def _finish(self, s, notes):
        _textbox(s, W - MX - 240, PAGE_Y - SIZE["page"] * LINE["page"], 240, SIZE["page"] * LINE["page"],
                 self._page[s.slide_id], SIZE["page"], MUTED, line=LINE["page"], align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.BOTTOM)
        self._notes(s, notes)

    def _notes(self, s, notes):
        if notes:
            _fill_frame(s.notes_slide.notes_text_frame, notes, SIZE["notes"], TEXT, line=1.4)

    def _check_text(self, page, title, subtitle, key):
        words = len(_plain(title).split())
        if words > 5:
            self._warn(page, f"제목이 {words} 단어 — 2–4 단어의 짧은 주제로")
        if _lines(title, SIZE["title"], CW, bold=True) > 1:
            self._warn(page, "제목이 한 줄을 넘음")
        for name, t, size, most, bold in (("부제", subtitle, SIZE["subtitle"], 2, False),
                                          ("빨간 줄", key, SIZE["key"], 1, True)):
            if not t:
                continue
            if SENTENCE_END.search(_plain(t).strip()):
                self._warn(page, f"{name}가 문장형 (~다) — 음슴체·단정형으로 (~사용, ~확보, ~이내)")
            if _lines(t, size, CW, bold) > most:
                self._warn(page, f"{name}가 {most} 줄을 넘음 — 줄이거나 노트로")

    def _figure(self, s, x, y, w, h, panel, page):
        """이름표 + 그림 + 캡션 한 칸. 그림은 칸 왼쪽에 붙이고 비율을 지킨다. 돌려주는 것: 그림 폭."""
        label, cap = panel.get("label", ""), panel.get("caption", "")
        iw, ih = _image_size(panel["image"])
        lh = SIZE["label"] * LINE["label"] if label else 0
        img_top = y + (lh + 12 if label else 0)
        cap_n = _lines(cap, SIZE["caption"], w)
        for _ in range(3):                                # 캡션 줄 수가 그림 폭에 따라 바뀌어 몇 번 맞춘다
            ch = cap_n * SIZE["caption"] * LINE["caption"]
            room = y + h - img_top - (ch + 12 if cap else 0)
            sc = min(w / iw, room / ih)
            pw, ph = iw * sc, ih * sc
            n = _lines(cap, SIZE["caption"], pw)
            if n == cap_n:
                break
            cap_n = n
        if label:
            _textbox(s, x, y, w, lh, label, SIZE["label"], INK, bold=True, line=LINE["label"])
        _picture(s, panel["image"], x, img_top, pw, ph, panel.get("alt", ""))
        if cap:
            if cap_n > 2:
                self._warn(page, f"캡션이 {cap_n} 줄 — 한두 줄로")
            if SENTENCE_END.search(_plain(cap).strip()):
                self._warn(page, "캡션이 문장형 (~다) — 음슴체·단정형으로")
            _textbox(s, x, img_top + ph + 12, pw, ch, cap, SIZE["caption"], GREY, line=LINE["caption"])
        return pw

    # 장 종류 ---------------------------------------------------------------
    def cover(self, title, subtitle="", kicker="", name="", affiliation="", date="", notes=""):
        """표지 — 작은 대문자 회색 줄 → 검정 가로선 → 큰 제목 → 회색 부제 → 이름·소속 → 맨 아래 날짜."""
        s = self._new_slide()
        items = []                                         # (높이, 그리는 함수)
        if kicker:
            kh = SIZE["cover_kicker"] * LINE["cover"]
            items.append((kh, lambda y: _textbox(s, MX, y, CW, kh, kicker, SIZE["cover_kicker"], MUTED, bold=True,
                                                 line=LINE["cover"], caps=True, spacing=200)))
        items.append((6, lambda y: _rect(s, MX, y, 326, 6, INK)))
        tn = _lines(title, SIZE["cover_title"], CW, bold=True)
        th = tn * SIZE["cover_title"] * LINE["cover_title"]
        items.append((th, lambda y: _textbox(s, MX, y, CW, th, title, SIZE["cover_title"], INK, bold=True,
                                             line=LINE["cover_title"])))
        if subtitle:
            sh = _lines(subtitle, SIZE["subtitle"], CW) * SIZE["subtitle"] * LINE["cover"]
            items.append((sh, lambda y: _textbox(s, MX, y, CW, sh, subtitle, SIZE["subtitle"], GREY, line=LINE["cover"])))
        if name or affiliation:
            nh = SIZE["cover_name"] * LINE["cover"]
            hh = nh * 2 + 8 if (name and affiliation) else nh

            def who(y):
                if name:
                    _textbox(s, MX, y, CW, nh, name, SIZE["cover_name"], INK, bold=True, line=LINE["cover"])
                if affiliation:
                    _textbox(s, MX, y + (nh + 8 if name else 0), CW, nh, affiliation, SIZE["cover_name"], MUTED,
                             line=LINE["cover"])
            items.append((hh, who))
        total = sum(h for h, _ in items) + 36 * (len(items) - 1)
        y = 128 + max(0, (H - 128 - 160 - total) / 2)
        for h, draw in items:
            draw(y)
            y += h + 36
        if date:
            dh = SIZE["page"] * LINE["page"]
            _textbox(s, MX, PAGE_Y - dh, CW, dh, date, SIZE["page"], MUTED, line=LINE["page"], anchor=MSO_ANCHOR.BOTTOM)
        if tn > 3:
            self._warn("표지", "제목이 세 줄을 넘음")
        self._notes(s, notes)

    def panels(self, title, subtitle="", panels=(), key="", notes=""):
        """그림 1–3 칸 — 칸마다 이름표 (영어), 그림, 캡션 한 줄 (한글). panels: [{"label","image","caption","alt"}]."""
        s, top, bottom = self._start(title, subtitle, key)
        n = len(panels)
        if not 1 <= n <= 3:
            self._warn(self._page[s.slide_id], f"그림 {n} 칸 — 1–3 칸으로")
        gap = 32 if n >= 3 else 64
        cw = (CW - gap * (n - 1)) / max(n, 1)
        for i, panel in enumerate(panels):
            self._figure(s, MX + i * (cw + gap), top, cw, bottom - top, panel, self._page[s.slide_id])
        self._finish(s, notes)

    def figure_bullets(self, title, subtitle="", image="", label="", caption="", bullets=(), key="", notes="",
                       alt="", fig_width=0.5):
        """왼쪽 그림 (본문 폭의 fig_width) + 오른쪽 불릿 3–5 개."""
        s, top, bottom = self._start(title, subtitle, key)
        pw = self._figure(s, MX, top, CW * fig_width, bottom - top,
                          {"image": image, "label": label, "caption": caption, "alt": alt}, self._page[s.slide_id])
        bx = MX + pw + 64
        bw = MX + CW - bx
        by = top + (SIZE["label"] * LINE["label"] + 12 if label else 0)
        need = (_lines(list(bullets), SIZE["body"], bw - SIZE["body"]) * SIZE["body"] * LINE["body"]
                + 16 * max(len(bullets) - 1, 0))
        if need > bottom - by:
            self._warn(self._page[s.slide_id], "불릿이 칸을 넘음 — 줄이거나 노트로")
        if len(bullets) > 5:
            self._warn(self._page[s.slide_id], f"불릿 {len(bullets)} 개 — 3–5 개로")
        _textbox(s, bx, by, bw, bottom - by, list(bullets), SIZE["body"], TEXT, line=LINE["body"], bullets=True,
                 gap_after=16)
        self._finish(s, notes)

    def _cards_row(self, s, x, y, w, max_h, cards, page):
        """카드를 (x, y) 부터 폭 w 에 나란히 — 높이는 가장 긴 카드에 맞추고 max_h 를 넘지 않는다."""
        n, gap, pad = len(cards), 32, 36
        cw = (w - gap * (n - 1)) / max(n, 1)
        inner = cw - 2 * pad
        th = SIZE["card_title"] * LINE["card"]

        def body_h(card):
            body = card.get("body", [])
            lines = _lines(list(body), SIZE["card_body"], inner - SIZE["card_body"])
            return lines * SIZE["card_body"] * LINE["card"] + 10 * max(len(body) - 1, 0)

        hgt = 2 * pad + th + 12 + max(body_h(c) for c in cards)
        if hgt > max_h:
            self._warn(page, "카드 글이 칸을 넘음 — 줄이거나 노트로")
            hgt = max_h
        for i, card in enumerate(cards):
            cx = x + i * (cw + gap)
            _rect(s, cx, y, cw, hgt, CARD_FILL, CARD_LINE, radius=8)
            _textbox(s, cx + pad, y + pad, inner, th, card["title"], SIZE["card_title"], INK, bold=True,
                     line=LINE["card"])
            _textbox(s, cx + pad, y + pad + th + 12, inner, hgt - 2 * pad - th - 12, list(card.get("body", [])),
                     SIZE["card_body"], TEXT, line=LINE["card"], bullets=True, gap_after=10)
        return hgt

    def cards(self, title, subtitle="", cards=(), key="", notes=""):
        """카드 2–3 개 — cards: [{"title": 영어 소제목, "body": [불릿, …]}]. 카드 높이는 가장 긴 카드에 맞춘다."""
        s, top, bottom = self._start(title, subtitle, key)
        self._cards_row(s, MX, top, CW, bottom - top, cards, self._page[s.slide_id])
        self._finish(s, notes)

    def equation(self, title, subtitle="", eq="", where="", image="", label="", caption="", cards=(), bullets=(),
                 key="", notes=""):
        """수식 장 — 위에 전체 폭 회색 수식 박스 (식은 가운데·원래 크기, 번호 없음, 아래 where 한 문장),
        그 아래 왼쪽 그림 + 오른쪽 카드 (또는 불릿). 그림이 없으면 카드·불릿이 전체 폭을 쓴다.
        eq: tools/equations.py 로 만든 PNG 경로, 또는 LaTeX 글 (여러 줄은 리스트, 맞출 자리에 &) — 글이면 그 자리에서 만든다."""
        s, top, bottom = self._start(title, subtitle, key)
        page = self._page[s.slide_id]
        png = _equation_png(eq)
        with Image.open(png) as im:
            (iw, ih), dpi = im.size, (im.info.get("dpi") or (600, 600))[0]
        ew, eh = iw / dpi * 144, ih / dpi * 144            # 그림의 실제 크기 (in) → 캔버스 px (1 in = 144 px)
        pad = 36
        room = CW - 2 * pad
        if ew > room:
            self._warn(page, "식이 박스보다 넓음 — 줄을 나눔")
            ew, eh = room, eh * room / ew
        wn = _lines(where, SIZE["card_body"], room) if where else 0
        if wn > 2:
            self._warn(page, "where 가 2 줄을 넘음 — 줄이거나 노트로")
        wh = wn * SIZE["card_body"] * LINE["card"]
        box_h = 2 * pad + eh + (16 + wh if where else 0)
        _rect(s, MX, top, CW, box_h, CARD_FILL, CARD_LINE, radius=8)
        _picture(s, png, MX + (CW - ew) / 2, top + pad, ew, eh, "equation")
        if where:
            _textbox(s, MX + pad, top + pad + eh + 16, room, wh, where, SIZE["card_body"], GREY, line=LINE["card"],
                     align=PP_ALIGN.CENTER)
        y = top + box_h + 40
        if (image or cards or bullets) and bottom - y < 120:
            self._warn(page, "수식 박스 아래 자리가 모자람 — 그림·카드를 줄이거나 장을 나눔")
        elif image or cards or bullets:
            x, w = MX, CW
            if image:
                pw = self._figure(s, MX, y, CW * 0.45, bottom - y, {"image": image, "label": label, "caption": caption},
                                  page)
                x = MX + pw + 64
                w = MX + CW - x
            if cards:
                self._cards_row(s, x, y, w, bottom - y, cards, page)
            elif bullets:
                _textbox(s, x, y, w, bottom - y, list(bullets), SIZE["body"], TEXT, line=LINE["body"], bullets=True,
                         gap_after=16)
        self._finish(s, notes)

    def table(self, title, subtitle="", header=(), rows=(), widths=None, key="", notes="", width=None):
        """표 — header: 머리 줄, rows: 줄 목록, widths: 칸 폭 비율. 이 연구가 나은 칸은 **굵게**."""
        s, top, bottom = self._start(title, subtitle, key)
        _, hgt = _table(s, MX, top, width or CW, header, rows, widths)
        if top + hgt > bottom:
            self._warn(self._page[s.slide_id], "표가 칸을 넘음 — 행을 줄이거나 장을 나눔")
        self._finish(s, notes)

    def summary(self, title, subtitle="", columns=(), closing="", notes=""):
        """요약 — 얇은 구분선 아래 3 단 열 (소제목 + 불릿), 맨 아래 회색 마무리 문장."""
        s, top, bottom = self._start(title, subtitle, "")
        _rule(s, MX, top, MX + CW)
        n, gap = len(columns), 64
        cw = (CW - gap * (n - 1)) / max(n, 1)
        th = SIZE["card_title"] * LINE["card"]
        ch = SIZE["body"] * LINE["body"] if closing else 0
        for i, col in enumerate(columns):
            x = MX + i * (cw + gap)
            _textbox(s, x, top + 32, cw, th, col["title"], SIZE["card_title"], INK, bold=True, line=LINE["card"])
            _textbox(s, x, top + 32 + th + 12, cw, bottom - top - 32 - th - 12 - ch - 40, list(col.get("body", [])),
                     SIZE["body"], TEXT, line=LINE["body"], bullets=True, gap_after=12)
        if closing:
            _textbox(s, MX, bottom - ch, CW, ch, closing, SIZE["body"], MUTED, line=LINE["body"])
        self._finish(s, notes)

    def appendix_index(self, rows, title="Appendix", subtitle="Question Index", notes="",
                       header=("질문", "Appendix", "장")):
        """Appendix 첫 장 (쪽 번호 A0) — rows: [(질문, "A1"), …]. 장 번호 칸은 save() 때 채운다. 8 줄이 넘으면 표 둘."""
        s, top, bottom = self._start(title, subtitle, "", page="A0")
        self._appendix = 0
        rows = list(rows)
        half = (len(rows) + 1) // 2
        groups = [rows] if len(rows) <= 8 else [rows[:half], rows[half:]]
        tw = (CW - 48) / 2
        for g, part in enumerate(groups):
            tbl, hgt = _table(s, MX + g * (tw + 48), top, tw, header, [(q, lab, "") for q, lab in part],
                              widths=(62, 22, 16))
            self._index_cells += [(tbl.cell(r, 2), lab) for r, (_, lab) in enumerate(part, start=1)]
            if top + hgt > bottom:
                self._warn("A0", "색인 표가 칸을 넘음")
        self._finish(s, notes)

    # 저장 ------------------------------------------------------------------
    def save(self, path):
        for cell, lab in self._index_cells:
            no = self._labels.get(lab)
            if no is None:
                self._warn("A0", f"색인의 {lab} 장이 덱에 없음")
            cell.text_frame.paragraphs[0].runs[0].text = str(no if no is not None else "?")
        cp = self.prs.core_properties                      # 작성자·회사 같은 개인 정보를 남기지 않는다
        cp.title = self.title
        cp.author = cp.last_modified_by = cp.subject = cp.keywords = cp.comments = cp.category = ""
        cp.revision = 1
        cp.created = cp.modified = datetime.datetime(2000, 1, 1)
        self.prs.save(str(path))
        _fixed_zip_dates(path)
        print(f"{Path(path).name}: {self._count} 장")
        for w in self.warnings:
            print("  경고 " + w)
        return self.warnings
